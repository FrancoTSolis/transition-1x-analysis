#!/usr/bin/env python3
"""Export the HTML deck to a 16:9 PowerPoint file.

Each slide is captured from index.html at 3x resolution and placed full-bleed
on a 13.333 x 7.5 in slide, so the .pptx looks exactly like the browser
version. The slides are pictures, not editable text: index.html /
template.html remain the source of truth for edits.

Run from gates_slides/:  python3 export_pptx.py  ->  quantum_ai_codesign.pptx
"""
from pathlib import Path

from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches

HERE = Path(__file__).resolve().parent
SHOTS = HERE / "art" / "_pptx"
OUT = HERE / "quantum_ai_codesign.pptx"
SCALE = 3


def capture() -> list[Path]:
    SHOTS.mkdir(parents=True, exist_ok=True)
    paths = []
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": 1320, "height": 780},
                                device_scale_factor=SCALE)
        page.goto((HERE / "index.html").as_uri())
        n = page.evaluate("document.querySelectorAll('.slide').length")
        for i in range(n):
            page.evaluate(f"showSlide({i})")
            # capture the slide box itself, without the page background
            page.evaluate("document.querySelector('.slide.active')"
                          ".style.transform = 'none'")
            page.wait_for_timeout(350)
            out = SHOTS / f"slide{i + 1}.png"
            page.locator(".slide.active").screenshot(path=str(out))
            paths.append(out)
            print(f"captured {out.name}")
        browser.close()
    return paths


def build(images: list[Path]) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for img in images:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(img), 0, 0,
                                 width=prs.slide_width, height=prs.slide_height)
    prs.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size / 1e6:.2f} MB, {len(images)} slides)")


if __name__ == "__main__":
    build(capture())
